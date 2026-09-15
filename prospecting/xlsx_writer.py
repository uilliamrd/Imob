"""Escrita da planilha (.xlsx) final de leads priorizados.

Mesmos dados do CSV, mas em formato de planilha real (cabeçalho fixo,
filtro automático, largura de coluna e uma escala de cor em score_final
para destacar visualmente as maiores oportunidades).
"""

from typing import List

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import OUTPUT_DIR
from .csv_writer import output_path
from .models import ROW_FIELDNAMES, LeadResult

COLUMN_WIDTHS = {
    "title": 32,
    "street": 40,
    "phone": 18,
    "categoryName": 22,
    "url": 45,
    "website": 32,
    "site_score": 11,
    "ads_score": 11,
    "score_final": 13,
    "motivo": 45,
}

SCORE_COLUMNS = {"site_score", "ads_score", "score_final"}


def write_xlsx(cidade: str, results: List[LeadResult], output_dir: str = OUTPUT_DIR) -> str:
    path = output_path(cidade, output_dir, "xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    ws.append(ROW_FIELDNAMES)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for r in results:
        row = r.to_row()
        ws.append([row[field] for field in ROW_FIELDNAMES])

    last_row = ws.max_row
    if last_row >= 2:
        last_col_letter = get_column_letter(len(ROW_FIELDNAMES))
        table_range = f"A1:{last_col_letter}{last_row}"

        table = Table(displayName="Leads", ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
        )
        ws.add_table(table)

        score_final_col = ROW_FIELDNAMES.index("score_final") + 1
        score_final_letter = get_column_letter(score_final_col)
        color_range = f"{score_final_letter}2:{score_final_letter}{last_row}"
        ws.conditional_formatting.add(
            color_range,
            ColorScaleRule(
                start_type="min", start_color="FFF8696B",
                mid_type="percentile", mid_value=50, mid_color="FFFFEB84",
                end_type="max", end_color="FF63BE7B",
            ),
        )

    for i, field in enumerate(ROW_FIELDNAMES, start=1):
        ws.column_dimensions[get_column_letter(i)].width = COLUMN_WIDTHS.get(field, 20)

    wb.save(path)
    return path
